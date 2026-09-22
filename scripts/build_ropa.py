#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Builds the RoPA (Record of Processing Activities) template as a 3-sheet
Excel workbook:
  1. RoPA Register   -- the fillable register itself, with one fully
                         worked example row (Row ID: EXAMPLE-01).
  2. Field Guide      -- field-by-field guidance notes AND the specific
                         DPDP / GDPR / CCPA statutory citation each field
                         exists to satisfy (audit-grade traceability --
                         see project README, Attack 6).
  3. Instructions      -- how to use this register, review cadence,
                         ownership, and the same DPDP commencement-status
                         caveat carried over from the compliance matrix.

Genericised from public statutory text only -- see README.md's
"Confidentiality" section. No client data, no client-specific
methodology, no proprietary deliverable structure of any kind.
"""
import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

# ============================================================
# SHEET 1 -- RoPA Register: columns + one fully worked example row
# ============================================================
REGISTER_COLUMNS = [
    "Record ID",
    "Processing Activity Name",
    "Business Function / Owning Team",
    "Data Fiduciary / Controller (Name + Contact)",
    "DPO / Named Contact for this Activity",
    "Purpose(s) of Processing",
    "Categories of Data Principals / Data Subjects",
    "Categories of Personal Data Processed",
    "Sensitive / Special-Category Data? (incl. children's data)",
    "Legal Basis / DPDP Ground",
    "Source of Collection",
    "Categories of Recipients",
    "Data Processors Engaged (Name, DPA in place?)",
    "Cross-Border Transfer? (Destination + Mechanism)",
    "Retention Period & Criteria",
    "Erasure Trigger Event",
    "Technical & Organisational Security Measures",
    "Automated Decision-Making / Profiling?",
    "Consent Manager Involved? (DPDP-specific)",
    "Significant Data Fiduciary Obligations Triggered?",
    "DPIA / Data Audit Conducted? (Date + Ref)",
    "Risk Rating",
    "Last Reviewed",
    "Next Review Due",
    "Record Owner",
]

EXAMPLE_ROW = [
    "EXAMPLE-01",
    "In-App User Behavioural Analytics & Personalised Recommendations",
    "Product / Growth (Analytics Team)",
    "[Fictional Co.] Pvt. Ltd. -- privacy@example.com",
    "J. Rao, Data Protection Officer -- dpo@example.com (India-based; "
        "appointment is precautionary, not yet SDF-mandated -- see col. T)",
    "To analyse in-app behavioural events (screens viewed, features used, "
        "session duration) in order to generate personalised content "
        "recommendations and measure feature adoption.",
    "Registered users of the [Fictional Co.] mobile app, aged 18+ "
        "(app is age-gated at sign-up; see col. I)",
    "Device ID, app usage/event logs, in-app navigation paths, "
        "self-reported content preferences, approximate location "
        "(city-level, derived from IP)",
    "No -- age-gating at sign-up is designed to exclude children's data "
        "from this activity entirely (see Common Implementation Pitfall "
        "in the DSAR/Breach templates re: not relying on self-declared age "
        "alone)",
    "Consent (DPDP S.4(1)(a)/S.6(1); GDPR Art 6(1)(a); CCPA notice-at-"
        "collection, no lawful-basis gate)",
    "Directly from the user (in-app event tracking, opt-in at first launch)",
    "Internal: Product, Growth, Data Science teams. External: the "
        "analytics-platform processor named below.",
    "Amplitude Analytics Inc. (US-based SaaS analytics processor) -- "
        "DPA executed 2025-11-01, incl. GDPR Art 28 processor terms and "
        "DPDP-equivalent processing-instruction clauses; DPA ref: "
        "LEGAL-DPA-0142",
    "Yes -- data processed by Amplitude Analytics Inc. on US-based "
        "infrastructure. Under DPDP S.16(1) (once in force -- see "
        "Instructions sheet), this destination has NOT been notified as "
        "restricted, so no DPDP-specific transfer mechanism is currently "
        "required; GDPR-side transfer (if any EU users) relies on SCCs "
        "executed within the same DPA (ref LEGAL-DPA-0142, Annex C)",
    "18 months from last user activity event, then aggregated/anonymised "
        "for trend reporting only",
    "Automatic purge job runs monthly against the 18-month rolling window; "
        "manual erasure also triggered by a verified deletion request "
        "(see DSAR Intake Procedure, Row ID mapping DSAR-ERASURE)",
    "Encryption in transit (TLS 1.2+) and at rest (AES-256); "
        "role-based access control restricted to Product/Growth/Data "
        "Science; quarterly access-review; event data pseudonymised "
        "(device ID, not account name) before reaching the analytics "
        "processor",
    "Yes -- personalised recommendation ranking is generated by a "
        "rules-plus-ML scoring model; no processing produces a legal or "
        "similarly significant effect on the user (recommendations are "
        "advisory/optional, not gating access to any service), so this "
        "sits outside GDPR Art 22's stricter automated-decision-making "
        "trigger, which is reserved for higher-stakes edge cases -- "
        "reassess if the model's output ever starts gating pricing, "
        "eligibility, or access",
    "No -- consent is captured directly in-app, not through a "
        "Board-registered Consent Manager",
    "Pending -- [Fictional Co.] has not been notified as a Significant "
        "Data Fiduciary under DPDP S.10(1); this record should be "
        "re-reviewed immediately upon any such notification (see "
        "compliance matrix Row SDF1)",
    "Internal privacy impact review conducted 2026-06-15 (ref: "
        "PIA-2026-014); formal DPIA under DPDP Rule 13 not applicable "
        "unless/until SDF-notified, and Rule 13 itself is not yet in "
        "force regardless (see Instructions sheet)",
    "Medium (profiling present, but no special-category data, no "
        "children's data, no legal/significant effect)",
    "2026-06-15",
    "2027-06-15 (or immediately upon SDF notification / a material "
        "change to the processing, whichever is sooner)",
    "Head of Product Analytics",
]

BLANK_ROW = ["" for _ in REGISTER_COLUMNS]
BLANK_ROW[0] = "[next Record ID]"

# ============================================================
# SHEET 2 -- Field Guide: guidance + statutory traceability per field
# ============================================================
FIELD_GUIDE_COLUMNS = [
    "Field",
    "Guidance",
    "DPDP Citation (obligation this field evidences)",
    "GDPR Citation",
    "CCPA/CPRA Citation",
]

FIELD_GUIDE_ROWS = [
    ("Record ID", "A stable, never-reused identifier for this processing activity. Do not recycle an ID after a record is retired -- keep a superseded record for audit-trail purposes instead of deleting it.",
     "S.8(4) (general accountability-programme evidence)", "Art 5(2) + Art 24(1) (demonstrability)", "No direct equivalent -- general recordkeeping best practice"),
    ("Processing Activity Name", "Plain-language name a non-specialist reviewer (e.g. a Board inquiry or a supervisory-authority auditor) could use to find this record without reading the full description.",
     "S.8(4)", "Art 30(1)(b)", "No direct equivalent"),
    ("Business Function / Owning Team", "The internal team accountable for this processing day-to-day -- not necessarily the same as the Record Owner (col. Y), who is accountable for keeping the RECORD accurate.",
     "S.8(4)", "Art 30(1)(a) implied (controller's own organisation)", "No direct equivalent"),
    ("Data Fiduciary / Controller (Name + Contact)", "The legal entity that determines the purpose and means of this processing. If more than one entity is a joint controller, name all of them and describe the arrangement.",
     "S.2(i) definition of Data Fiduciary", "Art 30(1)(a); Art 26 (joint controllers)", "Civil Code S.1798.140 definition of 'business'"),
    ("DPO / Named Contact for this Activity", "Every Data Fiduciary needs SOME named, reachable contact per S.8(9)/Rule 9 (once in force) -- a DPO specifically is only MANDATORY for a notified Significant Data Fiduciary (S.10(2)(a)). Do not assume this field requires a formally titled 'DPO' for every record.",
     "S.8(9) (general) / S.10(2)(a) (SDF-mandatory)", "Art 37-39", "No statutory DPO-equivalent role located"),
    ("Purpose(s) of Processing", "State the SPECIFIC purpose, not a generic category (\"analytics\" is too vague; \"measuring feature adoption to inform the product roadmap\" is specific). This is the anchor for testing purpose limitation later -- see compliance matrix Category 4.",
     "S.4(1)(a) read with S.6(1) ('specified purpose')", "Art 5(1)(b) + Art 30(1)(b)", "Civil Code S.1798.100(a)-(c) (notice at collection)"),
    ("Categories of Data Principals / Data Subjects", "Name the population, and flag explicitly if children could plausibly be included -- this single flag is what should trigger the DPDP S.9 children's-data controls (compliance matrix Rows CHD1-CHD2) elsewhere in your compliance programme.",
     "S.9 (if children are or could be included)", "Art 30(1)(c)", "Civil Code S.1798.120(c)-(d) (if minors included)"),
    ("Categories of Personal Data Processed", "List data CATEGORIES, not individual field names from your database schema -- the register should stay readable to a non-engineer reviewer.",
     "S.2(t) definition of personal data", "Art 30(1)(c); Art 4(1) definition", "Civil Code S.1798.140(v) definition of personal information"),
    ("Sensitive / Special-Category Data?", "DPDP itself does NOT define a 'sensitive personal data' special category the way GDPR Art 9 and CCPA's 'sensitive personal information' do (compliance matrix flags this structural difference where relevant) -- so this field is answered against the GDPR/CCPA definitions even for an India-only processing activity, as good practice.",
     "No DPDP special-category concept in the Act as enacted", "Art 9(1) (special categories)", "Civil Code S.1798.140(ae) (sensitive personal information)"),
    ("Legal Basis / DPDP Ground", "DPDP's architecture is binary and closed: consent (S.6) or one of the nine enumerated S.7 'certain legitimate uses.' Do not default to a GDPR-style 'legitimate interests' basis without checking it actually maps to a named S.7 category first -- see compliance matrix Row C1.",
     "S.4(1)(a)-(b) read with S.6 or S.7", "Art 6(1)(a)-(f)", "No lawful-basis gate in CCPA -- notice-at-collection duty only"),
    ("Source of Collection", "Direct-from-principal collection carries different notice obligations than third-party-sourced data. Flag third-party sources explicitly, and name the source.",
     "S.5 (notice) -- once Rule 3 is in force", "Art 14 (info where data not obtained from the subject)", "Civil Code S.1798.100(a)"),
    ("Categories of Recipients", "Internal teams AND external parties. Do not conflate this with the processor field below -- a recipient may be an internal team with no separate DPA.",
     "S.8(4) general obligations", "Art 30(1)(d)", "Civil Code S.1798.115 (categories disclosed to)"),
    ("Data Processors Engaged", "Name each processor and confirm a Data Processing Agreement is in place. DPDP does not use GDPR's controller/processor vocabulary as precisely, but the practical due-diligence expectation is the same.",
     "S.8(2) (Fiduciary remains responsible for processor compliance)", "Art 28 (processor)", "Civil Code S.1798.140(ag) (service provider)"),
    ("Cross-Border Transfer?", "State Yes/No, the destination country/territory, and the transfer mechanism relied on. IMPORTANT: DPDP's S.16 blocklist model means, as enacted, most destinations require NO specific transfer mechanism unless government-notified as restricted -- do not assume an SCC/BCR-equivalent is required by DPDP itself (see compliance matrix Row CB1's blocklist/allowlist finding). Still complete the GDPR-mechanism sub-field if any EU data subjects are in scope.",
     "S.16(1)-(2)", "Art 44-49 (Chapter V)", "No cross-border-specific mechanism in CCPA/CPRA"),
    ("Retention Period & Criteria", "State a period AND the criterion that starts the clock (e.g. 'from last activity,' not just a bare duration). DPDP's Rule 8 fixed-timeline/48-hour-warning mechanic is not yet in force (see Instructions sheet) -- until it is, S.8(7)-(8)'s bare 'necessary for the specified purpose' standard governs.",
     "S.8(7)-(8), read with Rule 8 (once in force)", "Art 5(1)(e)", "Civil Code S.1798.100(a)(3) (retention disclosure)"),
    ("Erasure Trigger Event", "Describe the actual mechanism (automatic job, manual process, or both) that enforces the retention period above -- a documented retention PERIOD with no enforcement mechanism is a common audit finding.",
     "S.8(7)-(8)", "Art 17", "Civil Code S.1798.105"),
    ("Technical & Organisational Security Measures", "Summarise, don't reproduce your full security policy here -- link out to it. DPDP's Rule 6 seven-part itemised minimum is not yet in force (see Instructions sheet); until it is, S.8(5)'s bare 'reasonable security safeguards' standard governs, so err toward GDPR Art 32's more developed standard as the de facto floor.",
     "S.8(5), read with Rule 6 (once in force)", "Art 32", "CCPA Regs Article 9 (if audit-threshold triggered)"),
    ("Automated Decision-Making / Profiling?", "State Yes/No, describe the logic in plain terms, and specifically assess whether the output produces a LEGAL or SIMILARLY SIGNIFICANT effect on the individual -- that's the GDPR Art 22 trigger, and DPDP has no equivalent provision at all (a genuine 'no equivalent' finding worth having ready).",
     "No DPDP provision located", "Art 22", "CCPA Regs Article 11 (ADMT) -- not independently verified for this project; flag for legal review, not asserted"),
    ("Consent Manager Involved?", "DPDP-specific field with no GDPR/CCPA counterpart -- see compliance matrix Row C6. Note Rule 4 (Consent Manager registration) and the underlying S.6(7)-(8) right are on DIFFERENT commencement dates (Nov 2026 vs May 2027 respectively).",
     "S.6(7)-(9), read with Rule 4 and the First Schedule", "No equivalent", "No equivalent"),
    ("Significant Data Fiduciary Obligations Triggered?", "SDF status is a discretionary GOVERNMENT NOTIFICATION event, not a self-assessed threshold -- do not mark 'No' just because your organisation doesn't meet GDPR's Art 37 or CCPA's Article 9 thresholds; conversely don't assume 'Yes' just because you do. See compliance matrix Row SDF1.",
     "S.10(1)(a)-(f)", "Art 37(1)(a)-(c) (different, self-assessed trigger)", "CCPA Regs S.7120(b) (different, self-assessed trigger)"),
    ("DPIA / Data Audit Conducted?", "A formal DPDP DPIA/independent-audit duty applies only to a notified SDF (S.10(2)(b), Rule 13) -- and Rule 13 itself is not yet in force. Record any INTERNAL privacy-impact review conducted as good practice regardless, as the worked example does.",
     "S.10(2)(b), read with Rule 13 (once in force)", "Art 35", "CCPA Regs Article 9-10 (if threshold triggered)"),
    ("Risk Rating", "Your own internal methodology -- this workbook does not prescribe one. Be consistent across records so the register can be sorted/triaged by risk.",
     "No specific statutory citation -- internal governance practice", "Art 35(1) risk-based approach (general principle)", "CCPA Regs S.7150 (risk-assessment triggers, if applicable)"),
    ("Last Reviewed / Next Review Due", "A record that is never reviewed is not evidence of an active accountability programme. Annual review is a reasonable floor; shorter for higher-risk records.",
     "S.8(4)", "Art 5(2) + Art 24(1)", "No direct equivalent"),
    ("Record Owner", "A named individual, not a team alias -- someone who can be asked, by name, why this record says what it says.",
     "S.8(4)", "Art 24(1)", "No direct equivalent"),
]

# ============================================================
# SHEET 3 -- Instructions
# ============================================================
INSTRUCTIONS_TEXT = [
    ("What this is",
     "A genericised Record of Processing Activities (RoPA) register template, structured to satisfy DPDP, GDPR, and CCPA/CPRA recordkeeping expectations in a single record per processing activity, rather than three separate registers. Built from public statutory text only -- see the repo README's Confidentiality section."),
    ("How to use it",
     "Duplicate the EXAMPLE-01 row for each new processing activity your organisation runs. Fill every column -- an intentionally blank field should say 'N/A' with a one-line reason, not be left empty, so a reviewer can tell 'not assessed' apart from 'assessed as not applicable.' Consult the Field Guide sheet for what each column is evidencing and its exact statutory basis."),
    ("IMPORTANT -- DPDP commencement status (as of this template's build date, 22 Sep 2026)",
     "Per Gazette Notification G.S.R. 843(E) (13 Nov 2025), DPDP Act Sections 3-17 -- essentially every substantive obligation this register documents -- do not come into force until 13 May 2027. Rules 3, 5-16, 22, 23 commence on the same date. This register documents the Act AS ENACTED, because that is the law a compliance programme has to build toward, but treat every DPDP citation in the Field Guide sheet as 'not yet legally binding, but the target state' until that date. See the compliance matrix's 'DPDP Commencement Status' column and the README's methodology section for the full finding and its verification trail."),
    ("Review cadence",
     "Review every record at least annually, and immediately upon: a material change to the processing; a new cross-border destination; onboarding a new processor; or (DPDP-specific) an SDF notification for your organisation."),
    ("Ownership",
     "Assign a single accountable owner for the register as a whole (typically the DPO or, absent a mandatory DPO, whoever holds the S.8(9)/Rule-9 named-contact role) in addition to each record's individual Record Owner."),
    ("Scope and disclaimer",
     "This is a practitioner reference tool for internal gap-assessment and scoping purposes. It is not legal advice and is not a substitute for qualified counsel. See the repo README for the full disclaimer."),
]

# ============================================================
# Build workbook
# ============================================================
def style_sheet(ws, col_widths, header_row=1, row_height=90):
    for col, w in col_widths.items():
        ws.column_dimensions[col].width = w
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row[0].row].height = row_height
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1).coordinate

register_df = pd.DataFrame([EXAMPLE_ROW, BLANK_ROW, BLANK_ROW, BLANK_ROW], columns=REGISTER_COLUMNS)
field_guide_df = pd.DataFrame(FIELD_GUIDE_ROWS, columns=FIELD_GUIDE_COLUMNS)
instructions_df = pd.DataFrame(INSTRUCTIONS_TEXT, columns=["Topic", "Detail"])

out_path = "../templates/RoPA_template.xlsx"
with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
    register_df.to_excel(writer, index=False, sheet_name="RoPA Register")
    field_guide_df.to_excel(writer, index=False, sheet_name="Field Guide")
    instructions_df.to_excel(writer, index=False, sheet_name="Instructions")

    ws1 = writer.sheets["RoPA Register"]
    widths1 = {get_column_letter(i + 1): 32 for i in range(len(REGISTER_COLUMNS))}
    widths1["A"] = 14
    style_sheet(ws1, widths1, row_height=140)
    # Highlight the worked example row distinctly
    for cell in ws1[2]:
        cell.fill = PatternFill("solid", fgColor="E2EFDA")

    ws2 = writer.sheets["Field Guide"]
    style_sheet(ws2, {"A": 34, "B": 70, "C": 45, "D": 40, "E": 45}, row_height=80)

    ws3 = writer.sheets["Instructions"]
    style_sheet(ws3, {"A": 40, "B": 100}, row_height=90)

print(f"Wrote RoPA template: {out_path} ({len(register_df)} register rows incl. 1 worked example, "
      f"{len(field_guide_df)} field-guide rows, {len(instructions_df)} instruction rows)")
